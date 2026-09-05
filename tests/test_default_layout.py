import os
import docx
import pytest
from app.services.ocr_service import sanitize_and_fix_paper_diagrams, get_sample_std6_maths_paper
from app.services.docx_generator import build_docx_paper

def test_paper_structure_generation_defaults_to_vertical():
    raw_paper = {
        'metadata': {
            'exam_title': 'Annual Exam',
            'standard': 'Std. 6th',
            'subject': 'Maths'
        },
        'sections': [
            {
                'id': 'sec_mcq',
                'title': 'Q-1 Choose the correct answer',
                'type': 'mcq',
                'questions': [
                    {
                        'text': '1) 3/4 + 7/4 = ______',
                        'options': ['(a) 10/4', '(b) 10/8', '(c) 2/6', '(d) 10/6']
                    }
                ]
            },
            {
                'id': 'sec_subq',
                'title': 'Q-2 Answer the following',
                'type': 'general',
                'questions': [
                    {'text': 'a) Solve 5 + 5'},
                    {'text': 'b) Solve 10 - 2'}
                ]
            }
        ]
    }

    sanitized = sanitize_and_fix_paper_diagrams(raw_paper)
    for sec in sanitized['sections']:
        assert sec.get('options_layout') == 'vertical'
        assert sec.get('subquestions_layout') == 'vertical'

    sample = get_sample_std6_maths_paper()
    for sec in sample['sections']:
        assert sec.get('options_layout') == 'vertical'
        # Clock section may use horizontal layout for side-by-side clock dials
        if not any('clock' in str(q.get('diagram_type', '')) for q in sec.get('questions', [])):
            assert sec.get('subquestions_layout') == 'vertical'

def test_mcq_docx_export_defaults_to_vertical_stacked(tmp_path):
    paper_data = {
        'metadata': {
            'exam_title': 'MCQ Exam',
            'standard': 'Std 6',
            'subject': 'Maths'
        },
        'sections': [
            {
                'id': 'sec_mcq',
                'title': 'Q-1 MCQ Questions',
                'type': 'mcq',
                # Neither options_layout nor subquestions_layout specified -> must default to vertical
                'questions': [
                    {
                        'text': '1) Which is a prime number?',
                        'options': ['(a) 4', '(b) 6', '(c) 7', '(d) 9']
                    }
                ]
            }
        ]
    }

    out_docx = str(tmp_path / "test_vertical_mcq.docx")
    build_docx_paper(paper_data, out_docx, temp_dir=str(tmp_path))

    assert os.path.exists(out_docx)
    doc = docx.Document(out_docx)

    # In vertical stacked layout, options are rendered as individual paragraphs rather than a horizontal table
    para_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert any('(a) 4' in t for t in para_texts)
    assert any('(b) 6' in t for t in para_texts)
    assert any('(c) 7' in t for t in para_texts)
    assert any('(d) 9' in t for t in para_texts)

    # Confirm there are no 4-column tables created for options
    for tbl in doc.tables:
        table_text = " ".join(c.text for row in tbl.rows for c in row.cells)
        assert '(a) 4' not in table_text
