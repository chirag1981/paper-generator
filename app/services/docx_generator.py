import os
import re
import logging
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.services.geometry_drawer import (
    generate_zigzag_diagram,
    generate_lines_rays_diagram,
    generate_pictograph_chart,
    generate_clock_diagram,
    generate_geometric_shape
)

logger = logging.getLogger(__name__)

# ==================== CONSTANTS & DESIGN TOKENS ====================
PRIMARY_COLOR = RGBColor(15, 23, 42)
MUTED_COLOR = RGBColor(71, 85, 105)
BORDER_COLOR = RGBColor(148, 163, 184)

# Project Standard Word Document Typographic Hierarchy & Spacing
DOCX_TITLE_FONT_SIZE = Pt(15)
DOCX_SECTION_HEADING_FONT_SIZE = Pt(13)
DOCX_BODY_FONT_SIZE = Pt(12)
DOCX_SUBTITLE_FONT_SIZE = Pt(12)

# Spacing standard tokens (Print-Ready Exam Paper Standards)
DOCX_LINE_SPACING = 1.15                        # Line spacing within multi-line questions (1.15-1.5)
DOCX_QUESTION_SPACE_AFTER = Pt(6)               # Consistent space after each question (tightened to prevent single-problem page spill)
DOCX_QUESTION_SPACING_BEFORE = Pt(4)            # Question space before (tightened to fit neatly into balanced pages)
DOCX_SECTION_SPACE_BEFORE = Pt(12)              # Space before new section heading (Q:1, Q:2, etc.)
DOCX_SECTION_SPACE_BEFORE_FIRST = Pt(6)         # Space before first section heading after header table
DOCX_HEADING_SPACE_AFTER = Pt(8)                # Option A: breathing room under heading (was effectively Pt(3) via 60 twips)
DOCX_SECTION_SPACE_AFTER = DOCX_HEADING_SPACE_AFTER

PAGE_CONTENT_WIDTH_INCHES = 7.2
PAGE_CONTENT_WIDTH = Inches(PAGE_CONTENT_WIDTH_INCHES)

DIAGRAM_WIDTHS = {
    'pictograph': {'side_by_side': Inches(2.9), 'standalone': Inches(3.4), 'inline': Inches(3.0)},
    'zigzag': {'side_by_side': Inches(2.2), 'standalone': Inches(2.2), 'inline': Inches(1.8)},
    'geometry_lines': {'side_by_side': Inches(2.7), 'standalone': Inches(2.8), 'inline': Inches(2.8)},
    'clock': {'side_by_side': Inches(1.5), 'standalone': Inches(1.8), 'inline': Inches(1.5)},
    'clock_blank': {'side_by_side': Inches(1.5), 'standalone': Inches(1.8), 'inline': Inches(1.5)},
    'cylinder': {'side_by_side': Inches(1.5), 'standalone': Inches(1.8), 'inline': Inches(1.3)},
    'pyramid': {'side_by_side': Inches(1.4), 'standalone': Inches(1.8), 'inline': Inches(1.2)},
    'sphere': {'side_by_side': Inches(1.3), 'standalone': Inches(1.6), 'inline': Inches(1.2)},
    'circle': {'side_by_side': Inches(1.3), 'standalone': Inches(1.6), 'inline': Inches(1.2)},
    'cone': {'side_by_side': Inches(1.3), 'standalone': Inches(1.6), 'inline': Inches(1.2)},
    'cube': {'side_by_side': Inches(1.3), 'standalone': Inches(1.6), 'inline': Inches(1.2)},
    'cuboid': {'side_by_side': Inches(1.3), 'standalone': Inches(1.6), 'inline': Inches(1.2)},
}

# Regex to match question numbering prefixes like 1., (1), 1), a., (a), A., (A), i., (i), Q.1, Q-1
QUESTION_PREFIX_REGEX = re.compile(
    r'^\s*(?:(?:Q|Que|Question|પ્રશ્ન)[\.\-\s]*[\d\u0AE6-\u0AEF]+[\.\:\)]*|[\(\[]\s*(?:[\d\u0AE6-\u0AEF]+|[a-zA-Z]|[ivxlcdmIVXLCDM]+)\s*[\)\]][\.\:]*|[\d\u0AE6-\u0AEF]+[\.\)\:]+|[a-zA-Z][\.\)]+|[ivxlcdmIVXLCDM]+[\.\)]+)\s*'
)


# ==================== XML & FORMATTING HELPERS ====================

def fmt_paragraph(p, before=Pt(0), after=Pt(0), spacing=1.0, keep_with_next=False):
    """Sets paragraph spacing and line spacing cleanly in a single call."""
    p.paragraph_format.space_before = before
    p.paragraph_format.space_after = after
    p.paragraph_format.line_spacing = spacing
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    return p


def set_cell_background(cell, hex_color: str):
    """Fills a table cell background with the given hex color string."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=35, bottom=35, left=60, right=60):
    """Sets internal padding (margins) in dxa for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color='cbd5e1', sz='4', val='single'):
    """Applies horizontal and inside borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def set_all_cell_borders(table, color='000000', sz='4', val='single'):
    """Sets complete box borders (including left and right outer borders) for grid tables."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def set_no_borders(table):
    """Removes all borders from a table for invisible layout grids."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def prevent_row_split(table):
    """Adds cantSplit to all rows in a table to prevent awkward page breaks across cells."""
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def set_repeat_header(table):
    """Sets repeat header row if table spans multiple pages."""
    if len(table.rows) > 0:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def apply_question_indent(paragraph, indent_inch: float = 0.3):
    """Applies clean hanging indent for questions starting with 1), a), (i), etc."""
    paragraph.paragraph_format.left_indent = Inches(indent_inch)
    paragraph.paragraph_format.first_line_indent = Inches(-indent_inch)


def set_run_font(run, font_name='Nirmala UI'):
    """Sets ASCII, High-ANSI, and Complex Script (Indic) fonts for clean rendering."""
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)


def create_omml_fraction(num: str, den: str, font_name: str = 'Cambria Math', is_bold: bool = False, font_size = DOCX_BODY_FONT_SIZE) -> str:
    """Generates an XML snippet for a native Microsoft Word stacked vertical fraction with font size scaling."""
    bold_xml = '<w:b/>' if is_bold else ''
    sz_pt = font_size.pt if hasattr(font_size, 'pt') else 12
    sz_val = int(sz_pt * 2)
    sz_xml = f'<w:sz w:val="{sz_val}"/><w:szCs w:val="{sz_val}"/>'
    return (
        f'<m:oMath {nsdecls("m")} {nsdecls("w")}>'
        f'<m:f>'
        f'<m:fPr><m:type m:val="bar"/></m:fPr>'
        f'<m:num><m:r><w:rPr>{bold_xml}{sz_xml}<w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}"/></w:rPr><m:t>{num}</m:t></m:r></m:num>'
        f'<m:den><m:r><w:rPr>{bold_xml}{sz_xml}<w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}"/></w:rPr><m:t>{den}</m:t></m:r></m:den>'
        f'</m:f>'
        f'</m:oMath>'
    )


def add_formatted_math_text(paragraph, text: str, font_name: str = 'Nirmala UI', font_size = DOCX_BODY_FONT_SIZE, is_bold: bool = False, color: RGBColor = PRIMARY_COLOR):
    """
    Parses text and inserts regular text runs and native Word OMML equations for stacked fractions.
    Recognizes patterns like 3/4, 7/4, 10/1000 and renders them as stacked vertical fractions.
    """
    frac_pattern = re.compile(r'(\d+)\s*/\s*(\d+)')
    last_end = 0

    for match in frac_pattern.finditer(text):
        start, end = match.span()
        if start > last_end:
            r = paragraph.add_run(text[last_end:start])
            set_run_font(r, font_name)
            r.font.size = font_size
            r.font.bold = is_bold
            r.font.color.rgb = color

        num, den = match.group(1), match.group(2)
        omml = create_omml_fraction(num, den, is_bold=is_bold, font_size=font_size)
        paragraph._p.append(parse_xml(omml))
        last_end = end

    if last_end < len(text):
        r = paragraph.add_run(text[last_end:])
        set_run_font(r, font_name)
        r.font.size = font_size
        r.font.bold = is_bold
        r.font.color.rgb = color


def render_question_text(paragraph, q_text: str, font_name: str = 'Nirmala UI', font_size = DOCX_BODY_FONT_SIZE, is_bold: bool = False, color: RGBColor = PRIMARY_COLOR, bold_number: bool = True):
    """
    Renders question text into a paragraph, ensuring the leading question number/identifier
    (e.g., '1.', '(1)', '1)', 'a.', '(a)', 'Q.1', '(i)') is rendered in BOLD, while the
    remaining question body follows the is_bold styling.
    """
    if not q_text:
        return

    m = QUESTION_PREFIX_REGEX.match(q_text)
    if bold_number and m:
        num_prefix = m.group(0)
        rem_text = q_text[m.end():]

        # Add number prefix in BOLD
        r_num = paragraph.add_run(num_prefix)
        set_run_font(r_num, font_name)
        r_num.font.size = font_size
        r_num.font.bold = True
        r_num.font.color.rgb = color

        # Add remaining question text
        if rem_text:
            add_formatted_math_text(paragraph, rem_text, font_name=font_name, font_size=font_size, is_bold=is_bold, color=color)
    else:
        add_formatted_math_text(paragraph, q_text, font_name=font_name, font_size=font_size, is_bold=is_bold, color=color)


def _safe_generate_diagram(diag_type: str, img_path: str) -> bool:
    """Safely generates diagram images with error handling to avoid breaking document build."""
    try:
        if diag_type == 'zigzag':
            generate_zigzag_diagram(img_path)
        elif diag_type == 'geometry_lines':
            generate_lines_rays_diagram(img_path)
        elif diag_type == 'pictograph':
            generate_pictograph_chart(img_path)
        elif diag_type in ['clock', 'clock_blank']:
            generate_clock_diagram(img_path, show_hands=(diag_type == 'clock'))
        elif diag_type in ['cylinder', 'pyramid', 'sphere', 'circle', 'cone', 'cube', 'cuboid']:
            generate_geometric_shape(img_path, diag_type)
        return os.path.exists(img_path)
    except Exception as e:
        logger.warning(f"Failed to generate {diag_type} diagram at '{img_path}': {e}")
        return False


def extract_question_summary(paper_data: dict) -> tuple:
    """
    Dynamically groups sections by their top-level Question (Q-1, Q-2, Q-3...)
    and calculates total marks allocated to each for the header assessment table.
    """
    custom_table = paper_data.get('custom_marks_table')
    if custom_table and isinstance(custom_table, list) and len(custom_table) > 0:
        q_list = []
        for item in custom_table:
            q_name = str(item.get('name', 'Q')).strip()
            try:
                q_marks = int(item.get('marks', 0))
            except (ValueError, TypeError):
                q_marks = 0
            q_list.append({'name': q_name, 'marks': q_marks})
        calc_total = sum(i['marks'] for i in q_list)
        return q_list, calc_total

    if paper_data.get('assessment_table_data'):
        data = paper_data['assessment_table_data']
        q_list = []
        for item in data.get('questions', []):
            q_name = item.get('name', 'Q')
            q_marks = int(item.get('marks', 0))
            q_list.append({'name': q_name, 'marks': q_marks})
        calc_total = sum(i['marks'] for i in q_list)
        return q_list, calc_total

    sections = paper_data.get('sections', [])
    if not sections:
        return ([{'name': 'Q-1', 'marks': 0}], 0)

    indic_to_eng = str.maketrans('૦૧૨૩૪૫૬૭૮૯०१२३४५६७८९', '01234567890123456789')
    q_dict = {}

    for idx, sec in enumerate(sections, 1):
        title = sec.get('title', '').strip()
        sec_id = str(sec.get('id', '')).lower()
        marks_str = str(sec.get('marks', '0')).translate(indic_to_eng)
        m_nums = re.findall(r'\d+', marks_str)
        marks_val = int(m_nums[0]) if m_nums else 0
        if marks_val == 0 and sec.get('questions'):
            marks_val = len(sec['questions'])

        clean_title = title.translate(indic_to_eng)
        match = re.search(r'(?:Q|Que|Question|Sec|Section|પ્ર|પ્રશ્ન|प्रश्न)[\.\s\-:_]*(\d+)', clean_title, re.IGNORECASE)
        if match:
            q_key = f"Q-{match.group(1)}"
        else:
            lead_digit = re.search(r'^[\s\-:_]*(\d+)', clean_title)
            if lead_digit:
                q_key = f"Q-{lead_digit.group(1)}"
            else:
                id_match = re.search(r'(?:sec|q)[\-_]*(\d+)', sec_id)
                if id_match:
                    q_key = f"Q-{id_match.group(1)}"
                else:
                    q_key = f"Q-{idx}"
                    logger.warning(f"Section title '{title}' did not match question pattern; grouped under fallback '{q_key}'")

        if q_key not in q_dict:
            q_dict[q_key] = 0
        q_dict[q_key] += marks_val

    q_list = [{'name': k, 'marks': v} for k, v in q_dict.items()]
    if not q_list:
        q_list = [{'name': 'Q-1', 'marks': 18}, {'name': 'Q-2', 'marks': 17}, {'name': 'Q-3', 'marks': 25}]

    calc_total = sum(item['marks'] for item in q_list)
    meta_total = str(paper_data.get('metadata', {}).get('total_marks', '')).translate(indic_to_eng)
    meta_nums = re.findall(r'\d+', meta_total)
    total_val = int(meta_nums[0]) if meta_nums else calc_total

    return q_list, total_val


# ==================== MODULAR SECTION RENDERERS ====================

def _build_header(doc, meta: dict, total_marks: str):
    """Renders the top exam header including Title, Subject, Standard, and Metadata table."""
    exam_title = (meta.get('exam_title') or 'Examination Paper').strip()
    standard = (meta.get('standard') or '').strip()
    subject = (meta.get('subject') or '').strip()
    date_val = meta.get('date', '')
    day_val = meta.get('day', '')
    roll_no_val = meta.get('roll_no', '')

    # 1. Exam Title (Centered, Bold)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_paragraph(title_p, before=Pt(0), after=Pt(0), spacing=1.0)
    run_title = title_p.add_run(exam_title)
    set_run_font(run_title, 'Nirmala UI')
    run_title.font.size = DOCX_TITLE_FONT_SIZE
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    # Blue bottom border below title
    pPr = title_p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="4" w:color="0284c7"/></w:pBdr>')
    pPr.append(pBdr)

    # 2. Subject (Centered, Bold - size 12)
    if subject:
        subj_p = doc.add_paragraph()
        subj_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(subj_p, before=Pt(0), after=Pt(0), spacing=1.0)
        prefix = 'Subject: ' if not any(subject.lower().startswith(p) for p in ['subject', 'વિષય', 'विषय']) else ''
        run_subj = subj_p.add_run(f'{prefix}{subject}')
        set_run_font(run_subj, 'Nirmala UI')
        run_subj.font.size = Pt(12)
        run_subj.font.bold = True
        run_subj.font.color.rgb = PRIMARY_COLOR

    # 3. Standard (Centered, Bold - size 12)
    if standard:
        std_p = doc.add_paragraph()
        std_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(std_p, before=Pt(0), after=Pt(0), spacing=1.0)
        run_std = std_p.add_run(f'{standard}')
        set_run_font(run_std, 'Nirmala UI')
        run_std.font.size = Pt(12)
        run_std.font.bold = True
        run_std.font.color.rgb = PRIMARY_COLOR

    # 4. Date, Day (Left) & Marks, Roll No (Right)
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_table.autofit = False
    meta_col_widths = [Inches(4.4), Inches(2.8)]
    for idx, col in enumerate(meta_table.columns):
        col.width = meta_col_widths[idx]
    for row in meta_table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = meta_col_widths[idx]
            set_cell_margins(cell, top=15, bottom=15, left=10, right=10)

    # Date / Marks
    m_r1c1 = meta_table.rows[0].cells[0].paragraphs[0]
    fmt_paragraph(m_r1c1, before=Pt(0), after=Pt(0), spacing=1.0)
    r = m_r1c1.add_run('Date: ')
    r.font.bold = True; r.font.size = DOCX_BODY_FONT_SIZE; set_run_font(r, 'Nirmala UI')
    r_d = m_r1c1.add_run(date_val if date_val else '____________')
    r_d.font.size = DOCX_BODY_FONT_SIZE; set_run_font(r_d, 'Nirmala UI')

    m_r1c2 = meta_table.rows[0].cells[1].paragraphs[0]
    m_r1c2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fmt_paragraph(m_r1c2, before=Pt(0), after=Pt(0), spacing=1.0)
    r = m_r1c2.add_run('Marks: ')
    r.font.bold = True; r.font.size = DOCX_BODY_FONT_SIZE; set_run_font(r, 'Nirmala UI')
    r_m = m_r1c2.add_run(str(total_marks))
    r_m.font.bold = True; r_m.font.size = DOCX_BODY_FONT_SIZE; set_run_font(r_m, 'Nirmala UI')

    # Day / Roll No
    m_r2c1 = meta_table.rows[1].cells[0].paragraphs[0]
    fmt_paragraph(m_r2c1, before=Pt(0), after=Pt(0), spacing=1.0)
    r = m_r2c1.add_run('Day: ')
    r.font.bold = True; r.font.size = DOCX_BODY_FONT_SIZE; set_run_font(r, 'Nirmala UI')
    r_dy = m_r2c1.add_run(day_val if day_val else '____________')
    r_dy.font.size = DOCX_BODY_FONT_SIZE; set_run_font(r_dy, 'Nirmala UI')

    m_r2c2 = meta_table.rows[1].cells[1].paragraphs[0]
    m_r2c2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fmt_paragraph(m_r2c2, before=Pt(0), after=Pt(0), spacing=1.0)
    r = m_r2c2.add_run('Roll No: ')
    r.font.bold = True; r.font.size = DOCX_BODY_FONT_SIZE; set_run_font(r, 'Nirmala UI')
    r_rn = m_r2c2.add_run(roll_no_val if roll_no_val else '____________')
    r_rn.font.size = DOCX_BODY_FONT_SIZE; set_run_font(r_rn, 'Nirmala UI')

    set_no_borders(meta_table)
    prevent_row_split(meta_table)


def _build_assessment_table(doc, q_summary: list, total_marks: str):
    """Builds the 3-row structured assessment marks and supervisor signature table."""
    num_q_cols = len(q_summary)
    total_cols = num_q_cols + 5  # Question, Q-1..Q-N, Total, Blank, Sig Label, Sig Box

    assessment_table = doc.add_table(rows=3, cols=total_cols)
    assessment_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    assessment_table.autofit = False
    set_all_cell_borders(assessment_table, color='000000', sz='4')
    prevent_row_split(assessment_table)

    col_w_question = Inches(1.05)
    col_w_total = Inches(0.6)
    col_w_blank = Inches(0.35)
    col_w_sig_lbl = Inches(1.1)
    col_w_sig_box = Inches(1.25)

    fixed_width = 1.05 + 0.6 + 0.35 + 1.1 + 1.25  # 4.35 inches
    avail_for_q = max(0.45 * num_q_cols, PAGE_CONTENT_WIDTH_INCHES - fixed_width)
    col_w_q = Inches(avail_for_q / max(1, num_q_cols))

    col_widths = [col_w_question] + [col_w_q] * num_q_cols + [col_w_total, col_w_blank, col_w_sig_lbl, col_w_sig_box]

    for idx, col in enumerate(assessment_table.columns):
        col.width = col_widths[idx]

    for r_idx in range(3):
        for c_idx in range(total_cols):
            cell = assessment_table.rows[r_idx].cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=35, bottom=35, left=35, right=35)

    blank_col_idx = num_q_cols + 2
    sig_lbl_col_idx = num_q_cols + 3
    total_col_idx = num_q_cols + 1

    # Merge the blank column vertically across rows 0, 1, 2
    merged_blank_cell = assessment_table.rows[0].cells[blank_col_idx].merge(assessment_table.rows[2].cells[blank_col_idx])
    merged_blank_cell.width = col_w_blank
    set_cell_margins(merged_blank_cell, top=35, bottom=35, left=15, right=15)

    # Row 0: "Question", Questions, "Total", [Blank], "Teacher's Sign", [Sig Box]
    p0 = assessment_table.rows[0].cells[0].paragraphs[0]
    fmt_paragraph(p0, before=Pt(0), after=Pt(0), spacing=1.0)
    r = p0.add_run('Question')
    r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')

    for i, q_item in enumerate(q_summary):
        p = assessment_table.rows[0].cells[i+1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(p, before=Pt(0), after=Pt(0), spacing=1.0)
        r = p.add_run(q_item['name'])
        r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')

    p_tot_hdr = assessment_table.rows[0].cells[total_col_idx].paragraphs[0]
    p_tot_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_paragraph(p_tot_hdr, before=Pt(0), after=Pt(0), spacing=1.0)
    r = p_tot_hdr.add_run('Total')
    r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')

    p_sig0 = assessment_table.rows[0].cells[sig_lbl_col_idx].paragraphs[0]
    fmt_paragraph(p_sig0, before=Pt(0), after=Pt(0), spacing=1.0)
    r = p_sig0.add_run("Teacher's Sign")
    r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')

    # Row 1: "Obtain Marks", blanks, blanks, [Blank], "Supe. Sign", [Sig Box]
    p1 = assessment_table.rows[1].cells[0].paragraphs[0]
    fmt_paragraph(p1, before=Pt(0), after=Pt(0), spacing=1.0)
    r = p1.add_run('Obtain Marks')
    r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')

    p_sig1 = assessment_table.rows[1].cells[sig_lbl_col_idx].paragraphs[0]
    fmt_paragraph(p_sig1, before=Pt(0), after=Pt(0), spacing=1.0)
    r = p_sig1.add_run("Supe. Sign")
    r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')

    # Row 2: "Marks", values, total, [Blank], "Rechk. Sign", [Sig Box]
    p2 = assessment_table.rows[2].cells[0].paragraphs[0]
    fmt_paragraph(p2, before=Pt(0), after=Pt(0), spacing=1.0)
    r = p2.add_run('Marks')
    r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')

    for i, q_item in enumerate(q_summary):
        p = assessment_table.rows[2].cells[i+1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(p, before=Pt(0), after=Pt(0), spacing=1.0)
        r = p.add_run(str(q_item['marks']))
        r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')

    p_tot_val = assessment_table.rows[2].cells[total_col_idx].paragraphs[0]
    p_tot_val.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_paragraph(p_tot_val, before=Pt(0), after=Pt(0), spacing=1.0)
    r = p_tot_val.add_run(str(total_marks))
    r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')

    p_sig2 = assessment_table.rows[2].cells[sig_lbl_col_idx].paragraphs[0]
    fmt_paragraph(p_sig2, before=Pt(0), after=Pt(0), spacing=1.0)
    r = p_sig2.add_run("Rechk. Sign")
    r.font.bold = True; r.font.size = Pt(10); set_run_font(r, 'Nirmala UI')


def _add_section_heading(doc, title_text: str, marks_text: str, is_first: bool = False):
    """Renders a section heading with right-aligned marks (e.g. [5], [10]) and 8pt breathing room under heading."""
    title_clean = (title_text or '').strip()
    marks_clean = (marks_text or '').strip()

    space_before = DOCX_SECTION_SPACE_BEFORE_FIRST if is_first else DOCX_SECTION_SPACE_BEFORE
    space_after = DOCX_HEADING_SPACE_AFTER

    p = doc.add_paragraph()
    fmt_paragraph(p, before=space_before, after=space_after, spacing=DOCX_LINE_SPACING, keep_with_next=True)
    p.paragraph_format.tab_stops.add_tab_stop(PAGE_CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)

    r0 = p.add_run(title_clean)
    set_run_font(r0, 'Nirmala UI')
    r0.font.size = DOCX_SECTION_HEADING_FONT_SIZE
    r0.font.bold = True
    r0.font.color.rgb = PRIMARY_COLOR

    if marks_clean:
        clean_marks = marks_clean.strip('[]() ')
        if clean_marks:
            clean_marks = f'[{clean_marks}]'
            p.add_run('\t')
            r1 = p.add_run(clean_marks)
            set_run_font(r1, 'Nirmala UI')
            r1.font.size = DOCX_SECTION_HEADING_FONT_SIZE
            r1.font.bold = True
            r1.font.color.rgb = PRIMARY_COLOR


def _render_options_box(doc, options_text: str):
    """Renders a dashed box with bracket options above questions."""
    opt_tbl = doc.add_table(rows=1, cols=1)
    opt_tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    opt_tbl.autofit = False
    opt_tbl.columns[0].width = PAGE_CONTENT_WIDTH
    opt_cell = opt_tbl.rows[0].cells[0]
    opt_cell.width = PAGE_CONTENT_WIDTH
    prevent_row_split(opt_tbl)
    set_cell_background(opt_cell, 'f1f5f9')
    set_cell_margins(opt_cell, top=50, bottom=50, left=80, right=80)

    opt_p = opt_cell.paragraphs[0]
    opt_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_paragraph(opt_p, before=Pt(0), after=Pt(4), spacing=1.0)
    r_opt = opt_p.add_run(options_text)
    set_run_font(r_opt, 'Nirmala UI')
    r_opt.font.size = DOCX_BODY_FONT_SIZE
    r_opt.font.bold = True
    r_opt.font.color.rgb = RGBColor(30, 41, 59)

    tblPr = opt_tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="4" w:space="0" w:color="64748b"/>'
        f'<w:bottom w:val="dashed" w:sz="4" w:space="0" w:color="64748b"/>'
        f'<w:left w:val="dashed" w:sz="4" w:space="0" w:color="64748b"/>'
        f'<w:right w:val="dashed" w:sz="4" w:space="0" w:color="64748b"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def _render_side_by_side_section(doc, sec: dict, questions: list, temp_dir: str):
    """Renders questions on the left and a diagram figure on the right in a 2-column layout."""
    diag_type = sec.get('diagram_type')
    diag_w = DIAGRAM_WIDTHS.get(diag_type, {}).get('side_by_side', Inches(2.7))
    diag_img_path = os.path.join(temp_dir, f"sec_{diag_type}.png")

    if not _safe_generate_diagram(diag_type, diag_img_path):
        diag_img_path = None

    sbs_tbl = doc.add_table(rows=1, cols=2)
    sbs_tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    sbs_tbl.autofit = False
    sbs_tbl.columns[0].width = Inches(4.3)
    sbs_tbl.columns[1].width = Inches(2.9)
    sbs_tbl.rows[0].cells[0].width = Inches(4.3)
    sbs_tbl.rows[0].cells[1].width = Inches(2.9)
    prevent_row_split(sbs_tbl)
    set_no_borders(sbs_tbl)

    cell_opts = sbs_tbl.rows[0].cells[0]
    cell_diag = sbs_tbl.rows[0].cells[1]
    set_cell_margins(cell_opts, top=10, bottom=10, left=5, right=15)
    set_cell_margins(cell_diag, top=10, bottom=10, left=10, right=5)

    # Left Cell: Options / Questions vertically
    for q_idx, q in enumerate(questions):
        p_q = cell_opts.paragraphs[0] if q_idx == 0 else cell_opts.add_paragraph()
        q_after = Pt(2) if q.get('answer_lines') else DOCX_QUESTION_SPACE_AFTER
        fmt_paragraph(p_q, before=DOCX_QUESTION_SPACING_BEFORE if q_idx > 0 else Pt(0), after=q_after, spacing=DOCX_LINE_SPACING, keep_with_next=bool(q.get('answer_lines')))

        if QUESTION_PREFIX_REGEX.match(q.get('text', '')):
            apply_question_indent(p_q, 0.3)
        render_question_text(p_q, q.get('text', ''), font_size=DOCX_BODY_FONT_SIZE, is_bold=q.get('is_bold', False))

        if q.get('answer_lines'):
            tot_lines = q.get('answer_lines', 1)
            q_txt = q.get('text', '')
            if '____' not in q_txt and '.......' not in q_txt:
                r_bl = p_q.add_run('   ______________________')
                set_run_font(r_bl, 'Nirmala UI')
                r_bl.font.size = Pt(9.5)
                r_bl.font.color.rgb = RGBColor(100, 116, 139)
            for line_idx in range(1, tot_lines):
                p_ans = cell_opts.add_paragraph()
                ans_after = DOCX_QUESTION_SPACE_AFTER if line_idx == tot_lines - 1 else Pt(2)
                fmt_paragraph(p_ans, before=Pt(1), after=ans_after, spacing=1.0, keep_with_next=(line_idx < tot_lines - 1))
                p_ans_r = p_ans.add_run('_________________________________________')
                set_run_font(p_ans_r, 'Nirmala UI')
                p_ans_r.font.size = Pt(9.5)
                p_ans_r.font.color.rgb = RGBColor(100, 116, 139)

    # Right Cell: Diagram picture
    if diag_img_path and os.path.exists(diag_img_path):
        p_diag = cell_diag.paragraphs[0]
        p_diag.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(p_diag, before=Pt(0), after=Pt(0), spacing=1.0)
        p_diag.add_run().add_picture(diag_img_path, width=diag_w)


def _render_standalone_diagram(doc, diag_type: str, temp_dir: str):
    """Renders a full-width centered diagram for stacked/standalone sections."""
    diag_w = DIAGRAM_WIDTHS.get(diag_type, {}).get('standalone', Inches(2.8))
    img_path = os.path.join(temp_dir, f"sec_{diag_type}.png")
    if _safe_generate_diagram(diag_type, img_path) and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(p_img, before=Pt(0), after=Pt(0), spacing=1.0)
        p_img.add_run().add_picture(img_path, width=diag_w)


def _render_table_data(doc, sec: dict):
    """Renders a bordered data table (e.g. for profit/loss, student items)."""
    tbl_info = sec.get('table_data')
    if not tbl_info:
        return
    headers = tbl_info.get('headers', [])
    rows = tbl_info.get('rows', [])
    if not headers or not rows:
        return

    t = doc.add_table(rows=len(rows)+1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    set_all_cell_borders(t, color='000000', sz='4')
    prevent_row_split(t)
    set_repeat_header(t)

    col_w = Inches(PAGE_CONTENT_WIDTH_INCHES / len(headers))
    for col in t.columns:
        col.width = col_w
    for h_i, h in enumerate(headers):
        cell = t.rows[0].cells[h_i]
        cell.width = col_w
        set_cell_background(cell, 'f1f5f9')
        set_cell_margins(cell, top=50, bottom=50, left=50, right=50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(p, before=Pt(0), after=Pt(0), spacing=1.0)
        add_formatted_math_text(p, h, font_size=DOCX_BODY_FONT_SIZE, is_bold=True)

    for r_i, r_data in enumerate(rows, start=1):
        for c_i in range(len(headers)):
            val = r_data[c_i] if c_i < len(r_data) else ''
            cell = t.rows[r_i].cells[c_i]
            cell.width = col_w
            set_cell_margins(cell, top=40, bottom=40, left=50, right=50)
            p = cell.paragraphs[0]
            fmt_paragraph(p, before=Pt(0), after=Pt(0), spacing=1.0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_i == 1 else WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_math_text(p, str(val), font_size=DOCX_BODY_FONT_SIZE)


def _is_row_tab_safe(items: list, num_columns: int = None) -> bool:
    """
    Checks whether a row of items can be safely rendered using tab-stops without
    overrunning column boundaries. Returns False if any item requires wrapping,
    contains diagrams, multi-line blanks, newlines, or exceeds the safe character budget.
    """
    if not items:
        return False
    if num_columns is None:
        num_columns = len(items)
    if num_columns < 2:
        return False

    col_w_inches = PAGE_CONTENT_WIDTH_INCHES / max(1, num_columns)
    # Safe character threshold: ~9.5 chars per inch at 12pt Nirmala UI font
    max_chars = max(10, int(col_w_inches * 9.5))

    shape_names = ['cylinder', 'pyramid', 'sphere', 'circle', 'cone', 'cube', 'cuboid']

    for item in items:
        if item is None:
            continue
        if isinstance(item, dict):
            if item.get('diagram_type'):
                return False
            if item.get('answer_lines', 0) > 1:
                return False
            if item.get('supercell_cells'):
                return False
            text = str(item.get('text', ''))
            q_lower = text.lower()
            if any(s in q_lower and ('shape' in q_lower or '[' in q_lower) for s in shape_names):
                return False
            has_single_blank = (item.get('answer_lines') == 1) and ('____' not in text) and ('.......' not in text)
        else:
            text = str(item)
            has_single_blank = False

        if '\n' in text or '\r' in text:
            return False

        frac_count = len(re.findall(r'\d+\s*/\s*\d+', text))
        effective_len = len(text.strip()) + (frac_count * 4) + (10 if has_single_blank else 0)

        if effective_len > max_chars:
            return False

    return True


def _populate_subquestion_cell(cell, item, col_width, font_size=DOCX_BODY_FONT_SIZE, bold: bool = False, temp_dir: str = None, sec_title: str = '', q_idx: int = 0, total_cols: int = 2):
    """Fills a single table cell with sub-question content (text, diagrams, blanks)."""
    cell.width = col_width
    set_cell_margins(cell, top=10, bottom=16, left=12, right=12)

    if not isinstance(item, dict):
        p_sub = cell.paragraphs[0]
        p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt_paragraph(p_sub, before=Pt(0), after=DOCX_QUESTION_SPACE_AFTER, spacing=DOCX_LINE_SPACING)
        render_question_text(p_sub, str(item), font_size=font_size, is_bold=bold, bold_number=True)
        return

    q = item
    q_diag = q.get('diagram_type')
    q_text = q.get('text', '')
    shape_diagrams = ['cylinder', 'pyramid', 'sphere', 'circle', 'cone', 'cube', 'cuboid']

    if not q_diag:
        q_lower = (q_text or '').lower()
        for s in shape_diagrams:
            if s in q_lower and ('shape' in q_lower or '[' in q_lower or 'shape' in (sec_title or '').lower()):
                q_diag = s
                break

    if q_diag in shape_diagrams:
        shape_path = os.path.join(temp_dir, f"q_{q_idx}_{q_diag}.png") if temp_dir else f"static/img/shape_{q_diag}.png"
        static_shape = os.path.join('static', 'img', f"shape_{q_diag}.png")
        if os.path.exists(static_shape):
            shape_path = static_shape
        elif not os.path.exists(shape_path):
            _safe_generate_diagram(q_diag, shape_path)

        p_sub = cell.paragraphs[0]
        p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt_paragraph(p_sub, before=Pt(2), after=DOCX_QUESTION_SPACE_AFTER if not q.get('answer_lines') else Pt(2), spacing=1.0)

        prefix_match = QUESTION_PREFIX_REGEX.match(q_text)
        num_prefix = prefix_match.group(0).strip() if prefix_match else f"({q_idx+1})"

        r_num = p_sub.add_run(f"{num_prefix}  ")
        set_run_font(r_num, 'Nirmala UI')
        r_num.font.size = font_size
        r_num.font.bold = True
        r_num.font.color.rgb = PRIMARY_COLOR

        img_w = Inches(0.95) if total_cols >= 3 else Inches(1.15)
        if os.path.exists(shape_path):
            p_sub.add_run().add_picture(shape_path, width=img_w)

        rem_text = QUESTION_PREFIX_REGEX.sub('', q_text).strip()
        clean_rem = re.sub(r'\[\s*(?:cylinder|pyramid|circle|sphere|cone|cube|cuboid)[^\]]*\]', '', rem_text, flags=re.IGNORECASE).strip()
        meaningful_text = re.sub(r'[=\._\s\-]+', '', clean_rem)

        if meaningful_text:
            r_txt = p_sub.add_run(f"  {clean_rem}")
            set_run_font(r_txt, 'Nirmala UI')
            r_txt.font.size = font_size
            r_txt.font.bold = True
        else:
            has_eq = '=' in rem_text or not rem_text
            eq_str = "=  " if has_eq else ""
            blank_len = "______________" if total_cols >= 3 else "____________________"
            r_line = p_sub.add_run(f"  {eq_str}{blank_len}")
            set_run_font(r_line, 'Nirmala UI')
            r_line.font.size = font_size
            r_line.font.bold = True
            r_line.font.color.rgb = RGBColor(71, 85, 105)

        if q.get('answer_lines') and q['answer_lines'] > 1:
            for l_idx in range(1, q['answer_lines']):
                p_ans = cell.add_paragraph()
                fmt_paragraph(p_ans, before=Pt(1), after=Pt(2), spacing=1.0)
                r_ans = p_ans.add_run('___________________________')
                set_run_font(r_ans, 'Nirmala UI')
                r_ans.font.size = Pt(9.5)
                r_ans.font.color.rgb = RGBColor(100, 116, 139)

    elif q_diag in ['clock', 'clock_blank']:
        clock_path = os.path.join(temp_dir, f"q_{q_idx}_{q_diag}.png") if temp_dir else f"static/img/{q_diag}.png"
        if not os.path.exists(clock_path):
            _safe_generate_diagram(q_diag, clock_path)
        p_clock = cell.paragraphs[0]
        p_clock.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(p_clock, before=Pt(2), after=Pt(4), spacing=1.0)
        img_clock_w = Inches(1.2) if total_cols >= 3 else Inches(1.5)
        if os.path.exists(clock_path):
            p_clock.add_run().add_picture(clock_path, width=img_clock_w)
        p_sub = cell.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(p_sub, before=Pt(0), after=DOCX_QUESTION_SPACE_AFTER if not q.get('answer_lines') else Pt(2), spacing=DOCX_LINE_SPACING)
        render_question_text(p_sub, q.get('text', ''), font_size=font_size, is_bold=q.get('is_bold', True), bold_number=True)
        if q.get('answer_lines'):
            tot_l = q['answer_lines']
            q_t = q.get('text', '')
            if '____' not in q_t and '.......' not in q_t:
                r_bl = p_sub.add_run('   ______________________')
                set_run_font(r_bl, 'Nirmala UI')
                r_bl.font.size = Pt(9.5)
                r_bl.font.color.rgb = RGBColor(100, 116, 139)
            for l_idx in range(1, tot_l):
                p_ans = cell.add_paragraph()
                ans_after = DOCX_QUESTION_SPACE_AFTER if l_idx == tot_l - 1 else Pt(1)
                fmt_paragraph(p_ans, before=Pt(1), after=ans_after, spacing=1.0)
                r_line = p_ans.add_run('___________________________')
                set_run_font(r_line, 'Nirmala UI')
                r_line.font.size = Pt(9.5)
                r_line.font.color.rgb = RGBColor(100, 116, 139)
    else:
        p_sub = cell.paragraphs[0]
        p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt_paragraph(p_sub, before=Pt(0), after=DOCX_QUESTION_SPACE_AFTER if not q.get('answer_lines') else Pt(2), spacing=DOCX_LINE_SPACING)
        render_question_text(p_sub, q.get('text', ''), font_size=font_size, is_bold=q.get('is_bold', bold), bold_number=True)
        if q.get('answer_lines'):
            tot_l = q['answer_lines']
            q_t = q.get('text', '')
            if '____' not in q_t and '.......' not in q_t:
                r_bl = p_sub.add_run('   ______________________')
                set_run_font(r_bl, 'Nirmala UI')
                r_bl.font.size = Pt(9.5)
                r_bl.font.color.rgb = RGBColor(100, 116, 139)
            for l_idx in range(1, tot_l):
                p_ans = cell.add_paragraph()
                ans_after = DOCX_QUESTION_SPACE_AFTER if l_idx == tot_l - 1 else Pt(1)
                fmt_paragraph(p_ans, before=Pt(1), after=ans_after, spacing=1.0)
                r_line = p_ans.add_run('___________________________')
                set_run_font(r_line, 'Nirmala UI')
                r_line.font.size = Pt(9.5)
                r_line.font.color.rgb = RGBColor(100, 116, 139)


def _render_tab_row(doc, items: list, font_size=DOCX_BODY_FONT_SIZE, bold: bool = False, num_columns: int = None):
    """Renders a row of short items as a single tab-aligned paragraph."""
    if not items:
        return
    if num_columns is None:
        num_columns = len(items)

    p = doc.add_paragraph()
    fmt_paragraph(p, before=Pt(0), after=DOCX_QUESTION_SPACE_AFTER, spacing=DOCX_LINE_SPACING)

    col_w = PAGE_CONTENT_WIDTH_INCHES / max(1, num_columns)
    for c_idx in range(1, num_columns):
        pos = Inches(c_idx * col_w)
        p.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.LEFT)

    for c_idx, item in enumerate(items):
        if c_idx > 0:
            p.add_run('\t')

        if isinstance(item, dict):
            text = item.get('text', '')
            item_bold = item.get('is_bold', bold)
            ans_lines = item.get('answer_lines', 0)
        else:
            text = str(item)
            item_bold = bold
            ans_lines = 0

        render_question_text(p, text, font_size=font_size, is_bold=item_bold, bold_number=True)
        if ans_lines == 1 and ('____' not in text) and ('.......' not in text):
            r_bl = p.add_run('   ______________________')
            set_run_font(r_bl, 'Nirmala UI')
            r_bl.font.size = Pt(9.5)
            r_bl.font.color.rgb = RGBColor(100, 116, 139)


def _render_table_row(doc, items: list, font_size=DOCX_BODY_FONT_SIZE, bold: bool = False, num_columns: int = None, temp_dir: str = None, sec_title: str = ''):
    """Renders a row of items as a 1-row borderless table with proper per-cell wrapping."""
    if not items:
        return
    if num_columns is None:
        num_columns = len(items)

    tbl = doc.add_table(rows=1, cols=num_columns)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    prevent_row_split(tbl)
    set_no_borders(tbl)

    col_width = Inches(PAGE_CONTENT_WIDTH_INCHES / max(1, num_columns))
    for col in tbl.columns:
        col.width = col_width

    for c_idx in range(num_columns):
        cell = tbl.rows[0].cells[c_idx]
        cell.width = col_width
        if c_idx < len(items):
            item = items[c_idx]
            _populate_subquestion_cell(cell, item, col_width, font_size=font_size, bold=bold, temp_dir=temp_dir, sec_title=sec_title, q_idx=c_idx, total_cols=num_columns)
        else:
            set_cell_margins(cell, top=10, bottom=16, left=12, right=12)


def render_row(doc, items: list, font_size=DOCX_BODY_FONT_SIZE, bold: bool = False, num_columns: int = None, temp_dir: str = None, sec_title: str = ''):
    """
    Renders a row of items (sub-questions or options).
    Content-aware: uses clean tab-stops if every item in the row is short and tab-safe;
    falls back to a borderless table with cell wrapping if any item is long or complex.
    """
    if not items:
        return
    if num_columns is None:
        num_columns = len(items)

    if _is_row_tab_safe(items, num_columns=num_columns):
        _render_tab_row(doc, items, font_size=font_size, bold=bold, num_columns=num_columns)
    else:
        _render_table_row(doc, items, font_size=font_size, bold=bold, num_columns=num_columns, temp_dir=temp_dir, sec_title=sec_title)


def add_single_line(doc, item, font_size=DOCX_BODY_FONT_SIZE, bold: bool = False, temp_dir: str = None, q_idx: int = 0):
    """Renders an odd trailing item as a single full-width line, not a lone half-width column."""
    if isinstance(item, dict):
        q_text = item.get('text', '')
        q_diag = item.get('diagram_type')
        q_bold = item.get('is_bold', bold)
        if q_diag or item.get('answer_lines') or item.get('supercell_cells'):
            _render_general_question(doc, item, q_text, temp_dir=temp_dir, q_idx=q_idx)
        else:
            p = doc.add_paragraph()
            fmt_paragraph(p, before=DOCX_QUESTION_SPACING_BEFORE if q_idx > 0 else Pt(0), after=DOCX_QUESTION_SPACE_AFTER, spacing=DOCX_LINE_SPACING)
            if QUESTION_PREFIX_REGEX.match(q_text):
                apply_question_indent(p, 0.3)
            render_question_text(p, q_text, font_size=font_size, is_bold=q_bold, bold_number=True)
    else:
        text = str(item)
        p = doc.add_paragraph()
        fmt_paragraph(p, before=DOCX_QUESTION_SPACING_BEFORE if q_idx > 0 else Pt(0), after=DOCX_QUESTION_SPACE_AFTER, spacing=DOCX_LINE_SPACING)
        if QUESTION_PREFIX_REGEX.match(text):
            apply_question_indent(p, 0.3)
        render_question_text(p, text, font_size=font_size, is_bold=bold, bold_number=True)


def render_two_column_grid(doc, questions: list, font_size=DOCX_BODY_FONT_SIZE, temp_dir: str = None, sec_title: str = ''):
    """
    Renders sub-questions in a 2-column grid. Each row of 2 items is
    independently checked for tab-safety — mixed short/long rows in the
    same grid render correctly rather than forcing one strategy for all.
    """
    if not questions:
        return

    for i in range(0, len(questions), 2):
        pair = questions[i:i+2]
        if isinstance(pair[0], dict):
            texts = [q.get('text', '') for q in pair]
            is_bold = pair[0].get('is_bold', False)
            has_rich_data = any(q.get('diagram_type') or q.get('answer_lines') or q.get('supercell_cells') for q in pair)
            items_to_render = pair if has_rich_data else texts
        else:
            texts = [str(q) for q in pair]
            is_bold = False
            items_to_render = texts

        if len(pair) == 2:
            render_row(doc, items_to_render, font_size=font_size, bold=is_bold, num_columns=2, temp_dir=temp_dir, sec_title=sec_title)
        else:
            # odd item out — render as a single full-width line, not a lone column
            odd_item = pair[0] if (isinstance(pair[0], dict) and (pair[0].get('diagram_type') or pair[0].get('answer_lines'))) else texts[0]
            add_single_line(doc, odd_item, font_size=font_size, bold=is_bold, temp_dir=temp_dir, q_idx=i)


def render_horizontal_row(doc, questions: list, font_size=DOCX_BODY_FONT_SIZE, bold: bool = False, temp_dir: str = None, sec_title: str = ''):
    """
    Renders sub-questions in a single horizontal row across N columns.
    Content-aware: uses clean tab-stops if short, falls back to borderless table if any item wraps.
    """
    if not questions:
        return
    is_bold = questions[0].get('is_bold', bold) if isinstance(questions[0], dict) else bold
    render_row(doc, questions, font_size=font_size, bold=is_bold, num_columns=len(questions), temp_dir=temp_dir, sec_title=sec_title)


def _render_horizontal_subquestions(doc, questions: list, layout: str = 'horizontal', temp_dir: str = None, sec_title: str = ''):
    """Renders short sub-questions in content-aware equal-width columns (1 row or 2 columns side-by-side)."""
    is_two_col = layout in ['two_columns', '2_columns', '2col', 'two-columns']
    if is_two_col:
        render_two_column_grid(doc, questions, font_size=DOCX_BODY_FONT_SIZE, temp_dir=temp_dir, sec_title=sec_title)
    else:
        render_horizontal_row(doc, questions, font_size=DOCX_BODY_FONT_SIZE, temp_dir=temp_dir, sec_title=sec_title)


def _render_mcq_question(doc, q: dict, q_text: str, q_idx: int = 0, default_layout: str = 'horizontal'):
    """Renders an MCQ question and its options in horizontal, 2-column side-by-side, or vertical layout."""
    p = doc.add_paragraph()
    fmt_paragraph(p, before=DOCX_QUESTION_SPACING_BEFORE if q_idx > 0 else Pt(0), after=Pt(3), spacing=DOCX_LINE_SPACING, keep_with_next=True)
    if QUESTION_PREFIX_REGEX.match(q_text):
        apply_question_indent(p, 0.3)
    render_question_text(p, q_text, font_size=DOCX_BODY_FONT_SIZE, is_bold=True, bold_number=True)

    opts = q.get('options', [])
    if not opts:
        return

    layout = q.get('options_layout') or default_layout or 'horizontal'
    is_two_col = layout in ['two_columns', '2_columns', '2col', 'two-columns']
    is_vertical = layout in ['vertical', 'stacked']

    if is_two_col:
        render_two_column_grid(doc, opts, font_size=DOCX_BODY_FONT_SIZE)
    elif is_vertical:
        for i, opt in enumerate(opts):
            p_opt = doc.add_paragraph()
            opt_after = DOCX_QUESTION_SPACE_AFTER if i == len(opts) - 1 else Pt(3)
            fmt_paragraph(p_opt, before=Pt(1), after=opt_after, spacing=1.0, keep_with_next=(i < len(opts) - 1))
            apply_question_indent(p_opt, 0.4)
            render_question_text(p_opt, opt, font_size=DOCX_BODY_FONT_SIZE, bold_number=True)
    else:
        # Default: horizontal 1 row
        render_row(doc, opts, font_size=DOCX_BODY_FONT_SIZE, bold=False, num_columns=len(opts))


def _render_true_false_question(doc, q_text: str, q_idx: int = 0):
    """Renders a True/False question with right-aligned bracket box."""
    tbl_tf = doc.add_table(rows=1, cols=2)
    tbl_tf.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_tf.autofit = False
    prevent_row_split(tbl_tf)
    tbl_tf.columns[0].width = Inches(6.0)
    tbl_tf.columns[1].width = Inches(1.2)
    tbl_tf.rows[0].cells[0].width = Inches(6.0)
    tbl_tf.rows[0].cells[1].width = Inches(1.2)
    set_no_borders(tbl_tf)

    set_cell_margins(tbl_tf.rows[0].cells[0], top=28, bottom=28, left=10, right=10)
    set_cell_margins(tbl_tf.rows[0].cells[1], top=28, bottom=28, left=10, right=10)

    p0 = tbl_tf.rows[0].cells[0].paragraphs[0]
    fmt_paragraph(p0, before=DOCX_QUESTION_SPACING_BEFORE if q_idx > 0 else Pt(0), after=DOCX_QUESTION_SPACE_AFTER, spacing=DOCX_LINE_SPACING)
    if QUESTION_PREFIX_REGEX.match(q_text):
        apply_question_indent(p0, 0.3)
    render_question_text(p0, q_text, font_size=DOCX_BODY_FONT_SIZE)

    p1 = tbl_tf.rows[0].cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fmt_paragraph(p1, before=DOCX_QUESTION_SPACING_BEFORE if q_idx > 0 else Pt(0), after=DOCX_QUESTION_SPACE_AFTER, spacing=DOCX_LINE_SPACING)
    r1 = p1.add_run('[                  ]')
    set_run_font(r1, 'Nirmala UI')
    r1.font.size = DOCX_BODY_FONT_SIZE; r1.font.bold = True
    r1.font.color.rgb = PRIMARY_COLOR


def _render_fill_in_blanks_question(doc, q_text: str, q_idx: int = 0):
    """Renders a fill-in-the-blanks question with clean spacing."""
    p = doc.add_paragraph()
    fmt_paragraph(p, before=DOCX_QUESTION_SPACING_BEFORE if q_idx > 0 else Pt(0), after=DOCX_QUESTION_SPACE_AFTER, spacing=DOCX_LINE_SPACING)
    if QUESTION_PREFIX_REGEX.match(q_text):
        apply_question_indent(p, 0.3)
    render_question_text(p, q_text, font_size=DOCX_BODY_FONT_SIZE)


def _render_general_question(doc, q: dict, q_text: str, temp_dir: str, q_idx: int = 0):
    """Renders general math, descriptive, construction, or calculation questions."""
    diag_type = q.get('diagram_type')
    if not diag_type:
        q_lower = (q_text or '').lower()
        for s in ['cylinder', 'pyramid', 'sphere', 'circle', 'cone', 'cube', 'cuboid']:
            if s in q_lower and ('shape' in q_lower or '[' in q_lower):
                diag_type = s
                break
    is_shape = diag_type in ['cylinder', 'pyramid', 'sphere', 'circle', 'cone', 'cube', 'cuboid']

    if is_shape:
        # Render geometric shape question side-by-side: (1) [Shape Image] = ________________________
        img_path = os.path.join(temp_dir, f"q_{q_idx}_{diag_type}.png") if temp_dir else f"static/img/shape_{diag_type}.png"
        static_shape = os.path.join('static', 'img', f"shape_{diag_type}.png")
        if os.path.exists(static_shape):
            img_path = static_shape
        elif not os.path.exists(img_path):
            _safe_generate_diagram(diag_type, img_path)

        tbl_shape = doc.add_table(rows=1, cols=2)
        tbl_shape.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl_shape.autofit = False
        prevent_row_split(tbl_shape)
        set_no_borders(tbl_shape)

        c_left = tbl_shape.rows[0].cells[0]
        c_right = tbl_shape.rows[0].cells[1]
        c_left.width = Inches(2.2)
        c_right.width = Inches(PAGE_CONTENT_WIDTH_INCHES - 2.2)
        set_cell_margins(c_left, top=8, bottom=14, left=10, right=10)
        set_cell_margins(c_right, top=8, bottom=14, left=10, right=10)

        # Extract question number like (1), 1., etc.
        prefix_match = QUESTION_PREFIX_REGEX.match(q_text)
        num_prefix = prefix_match.group(0).strip() if prefix_match else f"({q_idx+1})"

        p_l = c_left.paragraphs[0]
        p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt_paragraph(p_l, before=Pt(2), after=Pt(2), spacing=1.0)
        r_num = p_l.add_run(f"{num_prefix}  ")
        set_run_font(r_num, 'Nirmala UI')
        r_num.font.size = DOCX_BODY_FONT_SIZE; r_num.font.bold = True
        r_num.font.color.rgb = PRIMARY_COLOR

        if os.path.exists(img_path):
            p_l.add_run().add_picture(img_path, width=Inches(1.2))

        p_r = c_right.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt_paragraph(p_r, before=Pt(8), after=Pt(2), spacing=1.0)
        has_eq = '=' in q_text
        eq_str = "=  " if has_eq else ""
        add_formatted_math_text(p_r, f"{eq_str}____________________________________________", font_size=DOCX_BODY_FONT_SIZE, is_bold=True)
        return

    p = doc.add_paragraph()
    q_after = Pt(3) if q.get('answer_lines') else DOCX_QUESTION_SPACE_AFTER
    fmt_paragraph(p, before=DOCX_QUESTION_SPACING_BEFORE if q_idx > 0 else Pt(0), after=q_after, spacing=DOCX_LINE_SPACING, keep_with_next=bool(q.get('answer_lines')))
    if QUESTION_PREFIX_REGEX.match(q_text):
        apply_question_indent(p, 0.3)
    render_question_text(p, q_text, font_size=DOCX_BODY_FONT_SIZE, is_bold=q.get('is_bold', False))

    if diag_type:
        diag_w = DIAGRAM_WIDTHS.get(diag_type, {}).get('inline', Inches(2.0))
        img_path = os.path.join(temp_dir, f"inline_{diag_type}.png")
        if _safe_generate_diagram(diag_type, img_path) and os.path.exists(img_path):
            p_diag = doc.add_paragraph()
            p_diag.alignment = WD_ALIGN_PARAGRAPH.CENTER
            diag_after = Pt(3) if q.get('answer_lines') else DOCX_QUESTION_SPACE_AFTER
            fmt_paragraph(p_diag, before=Pt(2), after=diag_after, spacing=1.0, keep_with_next=bool(q.get('answer_lines')))
            p_diag.add_run().add_picture(img_path, width=diag_w)

    if q.get('supercell_cells'):
        cells_data = q['supercell_cells']
        tbl_sc = doc.add_table(rows=1, cols=len(cells_data))
        tbl_sc.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl_sc.autofit = False
        prevent_row_split(tbl_sc)
        set_table_borders(tbl_sc, color='334155', sz='6')
        col_w_sc = Inches(PAGE_CONTENT_WIDTH_INCHES / len(cells_data))
        for col in tbl_sc.columns:
            col.width = col_w_sc
        for i, val in enumerate(cells_data):
            cell = tbl_sc.rows[0].cells[i]
            cell.width = col_w_sc
            set_cell_background(cell, 'f8fafc')
            set_cell_margins(cell, top=20, bottom=20, left=10, right=10)
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt_paragraph(p_c, before=Pt(0), after=Pt(0), spacing=1.0)
            r = p_c.add_run(str(val))
            set_run_font(r, 'Nirmala UI')
            r.font.size = DOCX_BODY_FONT_SIZE; r.font.bold = True
            r.font.color.rgb = PRIMARY_COLOR

    if q.get('answer_lines'):
        total_lines = q.get('answer_lines', 1)
        has_inline_blank = ('____' in q_text) or ('.......' in q_text)

        # Place blank line directly next to the question on line 1
        if not has_inline_blank:
            r_blank = p.add_run('   _________________________________________')
            set_run_font(r_blank, 'Nirmala UI')
            r_blank.font.size = Pt(9.5)
            r_blank.font.color.rgb = RGBColor(100, 116, 139)

        # For additional answer lines (if answer_lines > 1), add clean blank lines below without Ans:
        for line_idx in range(1, total_lines):
            p_ans = doc.add_paragraph()
            ans_after = DOCX_QUESTION_SPACE_AFTER if line_idx == total_lines - 1 else Pt(3)
            fmt_paragraph(p_ans, before=Pt(2), after=ans_after, spacing=1.0, keep_with_next=(line_idx < total_lines - 1))
            apply_question_indent(p_ans, 0.3)
            r_line = p_ans.add_run('____________________________________________________________________')
            set_run_font(r_line, 'Nirmala UI')
            r_line.font.size = Pt(9.5)
            r_line.font.color.rgb = RGBColor(100, 116, 139)

def render_vertical_question(doc, q, q_idx: int = 0, temp_dir: str = None, q_type: str = None, default_layout: str = 'horizontal'):
    """Renders a single question vertically (one paragraph/line per question)."""
    if isinstance(q, dict):
        q_text = q.get('text', '')
        q_has_opts = bool(q.get('options'))
        if str(q_type).lower() == 'mcq' or q_has_opts:
            _render_mcq_question(doc, q, q_text, q_idx, default_layout=default_layout)
        elif q_type == 'true_false':
            _render_true_false_question(doc, q_text, q_idx)
        elif q_type == 'fill_in_blanks':
            _render_fill_in_blanks_question(doc, q_text, q_idx)
        else:
            _render_general_question(doc, q, q_text, temp_dir, q_idx)
    else:
        p = doc.add_paragraph()
        fmt_paragraph(p, before=DOCX_QUESTION_SPACING_BEFORE if q_idx > 0 else Pt(0), after=DOCX_QUESTION_SPACE_AFTER, spacing=DOCX_LINE_SPACING)
        if QUESTION_PREFIX_REGEX.match(str(q)):
            apply_question_indent(p, 0.3)
        render_question_text(p, str(q), font_size=DOCX_BODY_FONT_SIZE, is_bold=False, bold_number=True)


def _render_drawing_boxes(doc, sec: dict, temp_dir: str = None):
    """Renders dashed boxes for student construction/drawing answers, or clock faces if time-related."""
    raw_boxes = sec.get('drawing_boxes', [])
    boxes = [b for b in raw_boxes if not any(k in str(b).lower() for k in ["figure showing", "geometric figure", "points l, m, p", "points o, c, a"])]
    if not boxes:
        return

    sec_title = (sec.get('title') or '').lower()
    is_clock_drawing = (
        any(k in sec_title for k in ['draw hands', 'clock', 'time', 'times']) or
        any(any(t in str(b).lower() for t in ['quarter', 'half past', 'o\'clock', 'past', ':']) for b in boxes)
    )

    if is_clock_drawing:
        tbl_box = doc.add_table(rows=1, cols=len(boxes))
        tbl_box.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl_box.autofit = False
        prevent_row_split(tbl_box)
        set_no_borders(tbl_box)
        col_w_box = Inches(PAGE_CONTENT_WIDTH_INCHES / len(boxes))
        for col in tbl_box.columns:
            col.width = col_w_box

        clock_path = os.path.join(temp_dir, 'clock_blank.png') if temp_dir else 'static/img/clock_blank.png'
        if not os.path.exists(clock_path):
            _safe_generate_diagram('clock_blank', clock_path)

        for i, label in enumerate(boxes):
            c = tbl_box.rows[0].cells[i]
            c.width = col_w_box
            set_cell_margins(c, top=10, bottom=20, left=15, right=15)
            p_img = c.paragraphs[0]
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt_paragraph(p_img, before=Pt(2), after=Pt(4), spacing=1.0)
            if os.path.exists(clock_path):
                p_img.add_run().add_picture(clock_path, width=Inches(1.5))
            p_lbl = c.add_paragraph()
            p_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt_paragraph(p_lbl, before=Pt(2), after=DOCX_QUESTION_SPACE_AFTER, spacing=DOCX_LINE_SPACING)
            lbl_text = str(label)
            if '___' not in lbl_text and '...' not in lbl_text:
                lbl_text = f"{lbl_text} ____________"
            add_formatted_math_text(p_lbl, lbl_text, font_size=DOCX_BODY_FONT_SIZE, is_bold=True)
        return

    tbl_box = doc.add_table(rows=1, cols=len(boxes))
    tbl_box.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_box.autofit = False
    prevent_row_split(tbl_box)
    set_table_borders(tbl_box, color='94a3b8', sz='4', val='dashed')
    box_margin_vert = 480 if len(boxes) <= 3 else 520
    col_w_box = Inches(PAGE_CONTENT_WIDTH_INCHES / len(boxes))
    for col in tbl_box.columns:
        col.width = col_w_box
    for i, label in enumerate(boxes):
        c = tbl_box.rows[0].cells[i]
        c.width = col_w_box
        set_cell_background(c, 'ffffff')
        set_cell_margins(c, top=box_margin_vert, bottom=box_margin_vert, left=20, right=20)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_paragraph(p, before=Pt(0), after=DOCX_QUESTION_SPACE_AFTER, spacing=1.0)
        r = p.add_run(label)
        set_run_font(r, 'Nirmala UI')
        r.font.size = DOCX_BODY_FONT_SIZE; r.font.bold = True
        r.font.color.rgb = PRIMARY_COLOR


# ==================== MAIN BUILDER FUNCTION ====================

def build_docx_paper(paper_data: dict, output_path: str, temp_dir: str = None) -> str:
    """
    Main builder producing print-ready, professional examination papers in DOCX format.
    Ensures modularity, clean typography, dynamic question grouping, and robust error handling.
    """
    if not temp_dir:
        temp_dir = os.path.dirname(output_path) or '.'
    os.makedirs(temp_dir, exist_ok=True)
    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)

    doc = docx.Document()

    # Set default style font with Indic script fallback
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Nirmala UI'
    font.size = DOCX_BODY_FONT_SIZE
    font.color.rgb = PRIMARY_COLOR
    fmt_paragraph(style, before=Pt(0), after=Pt(0), spacing=1.0)

    # Set Margins (0.4 in top/bottom, 0.5 in left/right)
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    meta = paper_data.get('metadata', {})
    q_summary, calc_total = extract_question_summary(paper_data)
    total_marks = meta.get('total_marks', '')
    if not total_marks or total_marks == '60':
        total_marks = str(calc_total)

    # 1. Header & Title Block
    _build_header(doc, meta, total_marks)

    # 2. Assessment & Marks Table (Optional)
    if paper_data.get('show_marks_table', True) is not False and not paper_data.get('hide_marks_table', False):
        _build_assessment_table(doc, q_summary, total_marks)

    # 3. Sections & Questions Loop
    sections = paper_data.get('sections', [])
    for s_idx, sec in enumerate(sections):
        if sec.get('page_break_before', False) and s_idx > 0:
            doc.add_page_break()

        title = sec.get('title', f'Section {s_idx+1}')
        marks = sec.get('marks', '')
        if not marks and sec.get('questions'):
            marks = f"[{len(sec['questions'])}]"
        _add_section_heading(doc, title, marks, is_first=(s_idx == 0))

        if sec.get('options_box'):
            _render_options_box(doc, sec['options_box'])

        if sec.get('intro_text'):
            p_in = doc.add_paragraph()
            fmt_paragraph(p_in, before=Pt(2), after=Pt(4), spacing=DOCX_LINE_SPACING, keep_with_next=True)
            add_formatted_math_text(p_in, sec['intro_text'], font_size=DOCX_BODY_FONT_SIZE)

        questions = sec.get('questions', [])
        q_type = sec.get('type', 'general')
        diag_type = sec.get('diagram_type')
        is_side_by_side = bool(diag_type and len(questions) > 0 and sec.get('diagram_layout') != 'stacked')

        if is_side_by_side:
            _render_side_by_side_section(doc, sec, questions, temp_dir)
        else:
            if diag_type:
                _render_standalone_diagram(doc, diag_type, temp_dir)

        if sec.get('table_data'):
            _render_table_data(doc, sec)

        render_general_questions = not is_side_by_side
        if render_general_questions:
            sec_opt_layout = sec.get('options_layout') or sec.get('subquestions_layout') or 'horizontal'
            sec_sub_layout = sec.get('subquestions_layout') or sec.get('options_layout') or 'horizontal'

            is_explicit_two_col = sec_sub_layout in ['two_columns', '2_columns', '2col', 'two-columns']
            is_explicit_horiz = sec_sub_layout in ['horizontal', '1_row']
            is_explicit_vert = sec_sub_layout in ['vertical', 'stacked']

            # Check if non-MCQ questions in this section should be rendered horizontally or 2-column grid
            has_options_any = any(bool(q.get('options')) for q in questions)
            is_mcq_sec = (str(q_type).lower() == 'mcq') or has_options_any

            if is_mcq_sec:
                for q_idx, q in enumerate(questions):
                    q_text = q.get('text', '')
                    _render_mcq_question(doc, q, q_text, q_idx, default_layout=sec_opt_layout)
            else:
                layout = sec.get('subquestions_layout')
                if not layout:
                    if len(questions) > 1:
                        if is_explicit_two_col or any(q.get('diagram_type') in ['clock', 'clock_blank'] for q in questions):
                            layout = 'two_columns'
                        elif is_explicit_horiz or (all(QUESTION_PREFIX_REGEX.match(q.get('text', '').strip()) or len(q.get('text', '')) < 45 for q in questions) and not any(q.get('answer_lines') for q in questions) and len(questions) <= 4):
                            layout = 'horizontal'
                        else:
                            layout = 'vertical'
                    else:
                        layout = 'vertical'

                if layout in ['two_columns', '2_columns', '2col', 'two-columns']:
                    norm_layout = 'two_columns'
                elif layout in ['horizontal', '1_row']:
                    norm_layout = 'horizontal'
                else:
                    norm_layout = 'vertical'

                if norm_layout == 'vertical' or len(questions) <= 1:
                    for q_idx, q in enumerate(questions):
                        render_vertical_question(doc, q, q_idx=q_idx, temp_dir=temp_dir, q_type=q_type, default_layout=sec_opt_layout)
                elif norm_layout == 'horizontal':
                    is_b = questions[0].get('is_bold', False) if (questions and isinstance(questions[0], dict)) else False
                    render_row(doc, questions, font_size=DOCX_BODY_FONT_SIZE, bold=is_b, temp_dir=temp_dir, sec_title=title)
                elif norm_layout == 'two_columns':
                    render_two_column_grid(doc, questions, font_size=DOCX_BODY_FONT_SIZE, temp_dir=temp_dir, sec_title=title)

        if sec.get('drawing_boxes'):
            _render_drawing_boxes(doc, sec, temp_dir=temp_dir)

    doc.save(output_path)
    return output_path
